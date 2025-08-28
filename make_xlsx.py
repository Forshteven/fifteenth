from openpyxl.workbook import Workbook
from module_11_1 import data_conc, data_gmo

wb = Workbook()
wb.remove(wb.active)

sheet1 = wb.create_sheet('Данные по бетону')
sheet2 = wb.create_sheet('Данные по ГМО')


headers_conc = list(data_conc.keys())
for col_num, header in enumerate(headers_conc):
    sheet1.cell(row=col_num + 1, column=1).value = header
    sheet1.column_dimensions['A'].width = 45

start_dates_conc = []
finish_dates_conc = []
for start_date, finish_date in data_conc.values():
    start_dates_conc.append(start_date)
    finish_dates_conc.append(finish_date)

for col_num, header in enumerate(start_dates_conc):
    sheet1.cell(row=col_num + 1, column=2).value = header
    sheet1.column_dimensions['B'].width = 15

for col_num, header in enumerate(finish_dates_conc):
    sheet1.cell(row=col_num + 1, column=3).value = header
    sheet1.column_dimensions['C'].width = 15

headers_gmo = list(data_gmo.keys())
for col_num, header in enumerate(headers_gmo):
    sheet2.cell(row=col_num + 1, column=1).value = header
    sheet2.column_dimensions['A'].width = 70

start_dates_gmo = []
finish_dates_gmo = []
for start_date, finish_date in data_gmo.values():
    start_dates_gmo.append(start_date)
    finish_dates_gmo.append(finish_date)

for col_num, header in enumerate(start_dates_gmo):
    sheet2.cell(row=col_num + 1, column=2).value = header
    sheet2.column_dimensions['B'].width = 15

for col_num, header in enumerate(finish_dates_gmo):
    sheet2.cell(row=col_num + 1, column=3).value = header
    sheet2.column_dimensions['C'].width = 15

wb.save("data.xlsx")
